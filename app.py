"""ATS Resume Checker - Streamlit + Gemini Flash."""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# ----------------------------- Config ---------------------------------------
DEFAULT_MODEL = "gemini-2.5-flash"  # override with GEMINI_MODEL in secrets/env
MAX_CHARS = 20000  # safety cap on text sent to the model
MIN_CHARS = 150  # below this we assume the file is scanned / empty

# Weights used to compute the final ATS score from the category scores
WEIGHTS = {
    "keyword_match": 0.30,
    "content_quality": 0.25,
    "formatting": 0.20,
    "impact_metrics": 0.15,
    "structure": 0.10,
}
LABELS = {
    "keyword_match": "Keyword match",
    "content_quality": "Content quality",
    "formatting": "ATS-friendly formatting",
    "impact_metrics": "Impact & metrics",
    "structure": "Structure & sections",
}

SYSTEM_PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and \
professional resume coach. Evaluate the resume honestly and strictly; do not \
inflate scores. Base every statement ONLY on the resume text provided. Never \
invent experience, skills or numbers.

Score each category from 0 to 100:
- keyword_match: relevant skills/keywords for the target role (use the job \
description if given, otherwise infer the most likely target role from the resume)
- content_quality: clarity, grammar, specificity, relevance
- formatting: ATS parseability (standard headings, no tables, columns or graphics, \
consistent dates, contact info present)
- impact_metrics: use of action verbs and quantified achievements
- structure: presence and order of Contact, Summary, Experience, Education, Skills

Return ONLY valid JSON with exactly this shape:
{
  "target_role": "string",
  "scores": {
    "keyword_match": 0, "content_quality": 0, "formatting": 0,
    "impact_metrics": 0, "structure": 0
  },
  "summary": "2-3 sentence overall assessment",
  "strengths": ["string"],
  "missing_keywords": ["string"],
  "improvements": [
    {"section": "string", "priority": "High|Medium|Low",
     "issue": "string", "suggestion": "string"}
  ],
  "rewrite_examples": [
    {"original": "string copied from resume", "improved": "string"}
  ]
}
Give 5-8 improvements ordered by priority and 2-4 rewrite examples. In \
rewrite examples, if the improved bullet needs a number you do not know, use a \
placeholder like [X%] instead of making one up."""


# ----------------------------- File parsing ---------------------------------
def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a PDF, DOCX or TXT upload."""
    name = filename.lower()
    if name.endswith(".pdf"):
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                reader.decrypt("")
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception:
            raise ValueError(
                "Could not read this PDF. It may be corrupted or password protected."
            )
    elif name.endswith(".docx"):
        try:
            doc = Document(io.BytesIO(data))
        except Exception:
            raise ValueError("Could not read this DOCX file. It may be corrupted.")
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT file.")

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text


# ----------------------------- AI helpers -----------------------------------
def parse_json(raw: str) -> dict:
    """Parse JSON from a model response, tolerating code fences / extra text."""
    raw = (raw or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        start, end = raw.find("{"), raw.rfind("}")
        if start != -1 and end > start:
            return json.loads(raw[start : end + 1])
        raise


def _clamp(value) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return 0


def normalize(result: dict) -> dict:
    """Fill defaults, clamp scores and compute the weighted overall score."""
    scores_in = result.get("scores") or {}
    scores = {k: _clamp(scores_in.get(k)) for k in WEIGHTS}
    overall = _clamp(sum(scores[k] * w for k, w in WEIGHTS.items()))

    def as_list(key):
        value = result.get(key)
        return value if isinstance(value, list) else []

    return {
        "target_role": str(result.get("target_role") or "Not detected"),
        "scores": scores,
        "overall": overall,
        "summary": str(result.get("summary") or ""),
        "strengths": [str(s) for s in as_list("strengths")],
        "missing_keywords": [str(s) for s in as_list("missing_keywords")],
        "improvements": [i for i in as_list("improvements") if isinstance(i, dict)],
        "rewrite_examples": [
            r for r in as_list("rewrite_examples") if isinstance(r, dict)
        ],
    }


def analyze_resume(client, model: str, resume_text: str, job_desc: str = "") -> dict:
    """Send the resume to Gemini and return the normalized analysis."""
    prompt = f"RESUME:\n\"\"\"\n{resume_text[:MAX_CHARS]}\n\"\"\"\n\n"
    if job_desc.strip():
        prompt += f"TARGET JOB DESCRIPTION:\n\"\"\"\n{job_desc[:MAX_CHARS]}\n\"\"\"\n"
    else:
        prompt += "No job description provided; infer the target role from the resume.\n"

    response = client.models.generate_content(
        model=model,
        contents=prompt,
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_PROMPT,
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    return normalize(parse_json(response.text))


def build_report(result: dict) -> str:
    """Create a downloadable Markdown report."""
    lines = [
        "# ATS Resume Report",
        f"**Overall ATS score:** {result['overall']}/100",
        f"**Target role:** {result['target_role']}",
        "",
        result["summary"],
        "",
        "## Category scores",
    ]
    lines += [f"- {LABELS[k]}: {v}/100" for k, v in result["scores"].items()]
    lines += ["", "## Strengths"] + [f"- {s}" for s in result["strengths"]]
    lines += ["", "## Missing keywords"] + [f"- {s}" for s in result["missing_keywords"]]
    lines += ["", "## Improvements"]
    for i in result["improvements"]:
        lines.append(
            f"- **[{i.get('priority', '')}] {i.get('section', '')}** - "
            f"{i.get('issue', '')} -> {i.get('suggestion', '')}"
        )
    lines += ["", "## Rewrite examples"]
    for r in result["rewrite_examples"]:
        lines.append(f"- Before: {r.get('original', '')}\n  After: {r.get('improved', '')}")
    return "\n".join(lines)


def get_secret(name: str, default: str = "") -> str:
    try:
        return st.secrets[name]
    except Exception:
        return os.environ.get(name, default)


# ----------------------------- UI -------------------------------------------
def score_color(score: int) -> str:
    return "🟢" if score >= 75 else "🟡" if score >= 50 else "🔴"


def main():
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an ATS score and AI-powered improvements.")

    api_key = get_secret("GEMINI_API_KEY")
    model = get_secret("GEMINI_MODEL", DEFAULT_MODEL)

    with st.sidebar:
        st.header("Settings")
        if not api_key:
            api_key = st.text_input("Gemini API key", type="password")
            st.caption("Get a free key at aistudio.google.com/apikey")
        else:
            st.success("API key loaded")
        model = st.text_input("Model", value=model)

    col1, col2 = st.columns(2)
    with col1:
        uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    with col2:
        job_desc = st.text_area(
            "Job description (optional, improves keyword matching)", height=150
        )

    if not st.button("Analyze resume", type="primary", disabled=uploaded is None):
        return
    if not api_key:
        st.error("Please provide a Gemini API key in the sidebar.")
        return

    try:
        with st.spinner("Reading your resume..."):
            text = extract_text(uploaded.name, uploaded.getvalue())
        if len(text) < MIN_CHARS:
            st.error(
                "Could not read enough text. If this is a scanned/image PDF, "
                "export a text-based PDF or upload a DOCX instead."
            )
            return
        with st.spinner("Analyzing with Gemini..."):
            client = genai.Client(api_key=api_key)
            result = analyze_resume(client, model, text, job_desc)
    except ValueError as e:
        st.error(str(e))
        return
    except json.JSONDecodeError:
        st.error("The AI returned an unreadable response. Please try again.")
        return
    except Exception as e:  # network, quota, invalid key, etc.
        st.error(f"Something went wrong: {e}")
        return

    st.divider()
    left, right = st.columns([1, 2])
    with left:
        st.metric("Overall ATS score", f"{result['overall']}/100")
        st.progress(result["overall"] / 100)
        st.write(f"**Target role:** {result['target_role']}")
    with right:
        st.write(result["summary"])

    st.subheader("Score breakdown")
    cols = st.columns(len(result["scores"]))
    for col, (key, val) in zip(cols, result["scores"].items()):
        col.metric(LABELS[key], f"{score_color(val)} {val}")

    s_col, k_col = st.columns(2)
    with s_col:
        st.subheader("✅ Strengths")
        for s in result["strengths"] or ["No strengths returned."]:
            st.write(f"- {s}")
    with k_col:
        st.subheader("🔑 Missing keywords")
        if result["missing_keywords"]:
            st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))
        else:
            st.write("None detected.")

    st.subheader("🛠 Improvements")
    order = {"high": 0, "medium": 1, "low": 2}
    for imp in sorted(
        result["improvements"], key=lambda i: order.get(str(i.get("priority", "")).lower(), 3)
    ):
        with st.expander(f"[{imp.get('priority', '-')}] {imp.get('section', 'General')}"):
            st.write(f"**Issue:** {imp.get('issue', '')}")
            st.write(f"**Fix:** {imp.get('suggestion', '')}")

    if result["rewrite_examples"]:
        st.subheader("✍️ Rewrite examples")
        for r in result["rewrite_examples"]:
            st.markdown(f"**Before:** {r.get('original', '')}")
            st.markdown(f"**After:** {r.get('improved', '')}")
            st.write("")

    st.download_button(
        "Download report (.md)",
        build_report(result),
        file_name="ats_report.md",
        mime="text/markdown",
    )


if __name__ == "__main__":
    main()
